"""Minimal Markdown → DOCX converter for the generated reports (headings, paragraphs, bullet / numbered
lists, pipe tables, images, fenced code, **bold** and `code` inline). Chinese text uses an East-Asian font."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK_FONT = "PingFang SC"
LATIN_FONT = "Calibri"
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _set_fonts(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None, mono: bool = False):
    run.font.name = "Consolas" if mono else LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_inline(par, text: str, size: float = 10.5):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _set_fonts(par.add_run(tok[2:-2]), size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            _set_fonts(par.add_run(tok[1:-1]), size - 0.5, mono=True, color=RGBColor(0x8B, 0x1A, 0x1A))
        elif tok.startswith("[") and "](" in tok:
            label, url = tok[1:].split("](", 1)
            _set_fonts(par.add_run(label), size, color=RGBColor(0x1F, 0x4E, 0x9C))
            _set_fonts(par.add_run(f" ({url[:-1]})"), size - 1, color=RGBColor(0x7F, 0x7F, 0x7F))
        else:
            _set_fonts(par.add_run(tok), size)


def _shade(cell, hex_fill: str):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _table(doc, rows: list[list[str]]):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            _add_inline(par, r[j] if j < len(r) else "", size=8.5 if len(rows) > 12 or ncol > 6 else 9)
            if i == 0:
                _shade(cell, "E8EEF6")
                for run in par.runs:
                    run.font.bold = True
    doc.add_paragraph()


def markdown_to_docx(md_path: Path, docx_path: Path, title: str | None = None) -> Path:
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = LATIN_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(10.5)
    for lvl in range(1, 4):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = LATIN_FONT
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        hs.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        hs.font.size = Pt({1: 18, 2: 14, 3: 12}[lvl])

    i = 0
    table_buf: list[list[str]] = []
    code_buf: list[str] = []
    in_code = False

    def flush_table():
        nonlocal table_buf
        if table_buf:
            _table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                _set_fonts(par.add_run("\n".join(code_buf)), 8.5, mono=True)
                par.paragraph_format.left_indent = Cm(0.5)
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        flush_table()
        s = line.rstrip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and title:
                text = title if not text else text
            h = doc.add_heading("", level=level)
            _add_inline(h, text, size={1: 18, 2: 14, 3: 12}[level])
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            img = (Path(md_path).parent / m.group(2)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(16.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        m = re.match(r"^\*([^*].*)\*$", s)
        if m and s.startswith("*图"):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_fonts(par.add_run(m.group(1)), 9, color=RGBColor(0x55, 0x55, 0x55))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:   # explicit numbers (Word's "List Number" style would continue numbering across lists)
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.9)
            par.paragraph_format.space_after = Pt(4)
            _set_fonts(par.add_run(f"{m.group(1)}.  "), 10.5, bold=True)
            _add_inline(par, m.group(2))
            i += 1
            continue
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            _add_inline(par, m.group(1))
            i += 1
            continue
        # paragraph (merge soft-wrapped lines)
        buf = [s]
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#{1,3}\s|!\[|\||```|\d+\.\s|[-*]\s|\*图)", lines[i + 1].strip()):
            i += 1
            buf.append(lines[i].strip())
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(6)
        _add_inline(par, " ".join(buf))
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            for run in par.runs:
                run.font.italic = True
        i += 1
    flush_table()
    docx_path = Path(docx_path)
    doc.save(docx_path)
    return docx_path
